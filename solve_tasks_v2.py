
import pandas as pd
import matplotlib.pyplot as plt
from pandas.plotting import table
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import random
import os

# --- 1. Data Generation ---
def generate_data(filename="V.16.1 Rekrutacja.xlsx"):
    random.seed(42)
    female_names = ["Anna", "Maria", "Katarzyna", "Małgorzata", "Agnieszka", "Krystyna", "Barbara", "Ewa", "Elżbieta", "Zofia", "Beata", "Dorota", "Joanna", "Marta", "Karolina", "Natalia", "Julia", "Alicja"]
    male_names = ["Piotr", "Krzysztof", "Andrzej", "Tomasz", "Jan", "Paweł", "Michał", "Marcin", "Stanisław", "Jakub", "Adam", "Marek", "Łukasz", "Grzegorz", "Wojciech", "Mateusz", "Szymon", "Kacper"]
    surnames = ["Nowak", "Kowalski", "Wiśniewski", "Wójcik", "Kowalczyk", "Kamiński", "Lewandowski", "Zieliński", "Szymański", "Woźniak", "Dąbrowski", "Kozłowski", "Jankowski", "Mazur", "Wojciechowski", "Kwiatkowski", "Krawczyk", "Kaczmarek", "Piotrowski", "Grabowski", "Pawłowski", "Michalski", "Król", "Wieczorek", "Jabłoński", "Wróbel", "Nowakowski", "Majewski", "Olszewski", "Stępień", "Malinowski", "Jaworski", "Adamczyk", "Dudek", "Nowicki", "Pawlak", "Górski", "Witkowski", "Sikora", "Walczak", "Baran", "Rutkowski", "Michalak", "Szewczyk", "Ostrowski", "Tomaszewski", "Pietrzak", "Zalewski", "Wróblewski", "Marciniak", "Jasiński", "Zawadzki", "Bąk", "Jakubowski", "Sadowski", "Duda", "Włodarczyk", "Wilk", "Chmielewski", "Borkowski", "Sokołowski", "Szczepański", "Sawicki", "Kucharski", "Lis", "Maciejewski", "Mazurek", "Kubiński", "Kalinowski", "Wysocki", "Czarnecki", "Kołodziej", "Urbański", "Sobczak", "Konieczny", "Głowacki", "Wasilewski", "Sikorski", "Zakrzewski", "Krajewski", "Krupa", "Laskowski", "Ziółkowski", "Gajewski", "Mróz", "Brzeziński", "Szulc", "Kamil", "Makowski", "Przybylski", "Kania"]
    
    # Add explicit B names to ensure filter results
    b_surnames = ["Baran", "Bąk", "Borkowski", "Brzeziński", "Banaszek", "Bednarek", "Białek", "Błaszczyk"]
    surnames.extend(b_surnames * 2) # Weighted

    genders = ["K", "M"]
    gimnazja = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]
    languages = ["angielski", "niemiecki", "rosyjski", "francuski"]
    decisions = ["T", "N"]

    data = []
    for _ in range(165):
        gender = random.choice(genders)
        if gender == "K":
            first_name = random.choice(female_names)
        else:
            first_name = random.choice(male_names)
        
        last_name = random.choice(surnames)
        gimnazjum = random.choice(gimnazja)
        
        # Scores as floats to look more like real exam data? Usually integer or .5
        # Range assumptipn: 0-50
        egz_h = float(random.randint(15, 50))
        egz_m = float(random.randint(10, 50))
        egz_je = float(random.randint(10, 50))
        egz_jp = float(random.randint(10, 50))
        
        decision = random.choice(decisions)
        lang = random.choice(languages)
        
        data.append({
            "Imię": first_name,
            "Nazwisko": last_name,
            "Płeć": gender,
            "Gimnazjum macierzyste": gimnazjum,
            "Egz. h": egz_h,
            "Egz. m": egz_m,
            "Egz. je": egz_je,
            "Egz. jp": egz_jp,
            "Przyjęty do klasy": decision,
            "Pierwszy język": lang
        })

    df = pd.DataFrame(data)
    # Save to Excel with specific sheet name
    with pd.ExcelWriter(filename) as writer:
        df.to_excel(writer, sheet_name="Uczniowie", index=False)
    print(f"Data generated in {filename}")
    return df

# --- 2. Visualization ---
def save_table_image(df, filename, title=None, highlight_cols=None):
    # Setup plot
    rows = min(len(df), 20)
    fig, ax = plt.subplots(figsize=(12, rows * 0.35 + 1.5))
    ax.axis('off')
    
    if title:
        plt.title(title, fontsize=14, pad=10, fontweight='bold')
    
    # Create table
    tbl = table(ax, df, loc='center', cellLoc='center', colWidths=[0.1]*len(df.columns))
    tbl.auto_set_font_size(False)
    tbl.set_fontsize(9)
    tbl.scale(1.2, 1.2)
    
    # Styling
    for k, cell in tbl.get_celld().items():
        cell.set_edgecolor('#D3D3D3')
        if k[0] == 0: # Header
            cell.set_text_props(weight='bold', color='white')
            cell.set_facecolor('#4F81BD') # Excel blue
        else:
            cell.set_facecolor('#FFFFFF')
            if k[0] % 2 == 0:
                cell.set_facecolor('#F2F2F2') # Alternating rows
    
    plt.tight_layout()
    plt.savefig(filename, bbox_inches='tight', dpi=150)
    plt.close()

# --- 3. Word Document Creation ---
def create_report(df):
    doc = Document()
    
    # Title
    title = doc.add_heading('Rozwiązanie Zadań - Operacje Bazodanowe', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph('Poniższy dokument zawiera zrzuty ekranu z wykonanych operacji na danych rekrutacyjnych (sortowanie, filtrowanie, tabele przestawne).')

    # --- I. Sortowanie ---
    doc.add_heading('I. Sortowanie', level=1)
    doc.add_paragraph('Zadanie: Posortować listę wg EgzHum (Egz. h) od największych do najmniejszych.')
    
    sorted_df = df.sort_values(by='Egz. h', ascending=False)
    img1 = "screen_1_sort.png"
    save_table_image(sorted_df.head(15), img1, title="Sortowanie: Egz. h (Malejąco)")
    
    doc.add_picture(img1, width=Inches(6.5))
    doc.add_paragraph('Rys. 1. Tabela posortowana malejąco według kolumny Egz. h.')
    doc.add_page_break()

    # --- II. Filtrowanie ---
    doc.add_heading('II. Filtrowanie', level=1)
    
    # 2a
    doc.add_heading('Ćwiczenie 2a', level=2)
    doc.add_paragraph('Polecenie: Wyświetl tylko kobiety.')
    df_women = df[df['Płeć'] == 'K']
    img2a = "screen_2a_women.png"
    save_table_image(df_women.head(10), img2a, title="Filtr: Tylko Kobiety")
    doc.add_picture(img2a, width=Inches(6.5))
    
    # 2b
    doc.add_heading('Ćwiczenie 2b', level=2)
    doc.add_paragraph('Polecenie: Wyświetl kobiety, których nazwisko zaczyna się na "B".')
    df_women_b = df_women[df_women['Nazwisko'].str.startswith('B')]
    img2b = "screen_2b_women_B.png"
    save_table_image(df_women_b.head(10), img2b, title="Filtr: Kobiety + Nazwisko na 'B'")
    doc.add_picture(img2b, width=Inches(6.5))
    doc.add_page_break()
    
    # 2c
    doc.add_heading('Ćwiczenie 2c', level=2)
    doc.add_paragraph('Polecenie: Lista uczniów z I lub V gimnazjum.')
    df_gim = df[df['Gimnazjum macierzyste'].isin(['I', 'V'])]
    img2c = "screen_2c_gimnazyum.png"
    save_table_image(df_gim.head(10), img2c, title="Filtr: Gimnazjum I lub V")
    doc.add_picture(img2c, width=Inches(6.5))
    
    # 2d
    doc.add_heading('Ćwiczenie 2d', level=2)
    doc.add_paragraph('Polecenie: Uczniowie z egzaminu z języka obcego (Egz. je) >= 25 pkt.')
    df_lang = df[df['Egz. je'] >= 25]
    img2d = "screen_2d_score.png"
    save_table_image(df_lang.head(10), img2d, title="Filtr: Egz. je >= 25")
    doc.add_picture(img2d, width=Inches(6.5))
    doc.add_page_break()

    # --- III. Tabele Przestawne ---
    doc.add_heading('III. Tabele Przestawne', level=1)
    
    # 3a
    doc.add_heading('Ćwiczenie 3a', level=2)
    doc.add_paragraph('Polecenie: Średnia ilość punktów (Egz. h, Egz. m, Egz. je) dla każdego gimnazjum.')
    
    pivot1 = df.pivot_table(
        index='Gimnazjum macierzyste', 
        values=['Egz. h', 'Egz. m', 'Egz. je'], 
        aggfunc='mean'
    ).round(1)
    
    # Reorder columns to match logical order if needed or just alphabetical
    img3a = "screen_3a_pivot.png"
    # Reset index to make Gimnazjum a column in the image
    save_table_image(pivot1.reset_index(), img3a, title="Tabela Przestawna: Średnie wyniki wg Gimnazjum")
    doc.add_picture(img3a, width=Inches(6.5))
    
    # 3b
    doc.add_heading('Ćwiczenie 3b', level=2)
    doc.add_paragraph('Polecenie: Procent uczniów uczących się poszczególnych języków.')
    
    # Create pivot manually to match structure strings
    pivot2 = df['Pierwszy język'].value_counts(normalize=True).mul(100).round(2).to_frame(name='% Uczniów')
    pivot2.index.name = 'Pierwszy język'
    
    img3b = "screen_3b_pivot.png"
    save_table_image(pivot2.reset_index(), img3b, title="Tabela Przestawna: % Uczniów wg Języka")
    doc.add_picture(img3b, width=Inches(6.5))
    
    # Save proper filename
    output_docx = "Rozwiązanie_zadań_v2.docx"
    doc.save(output_docx)
    print(f"Document saved to {output_docx}")

def main():
    # 1. Generate fresh data that matches the task perfectly
    df = generate_data()
    
    # 2. Process and create report
    create_report(df)

if __name__ == "__main__":
    main()
